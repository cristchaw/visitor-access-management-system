import csv

from datetime import timedelta

from django.http import HttpResponse
from openpyxl import Workbook

from django.db import models, transaction
from django.contrib import messages
from django.contrib.auth.decorators import login_required, user_passes_test
from django.shortcuts import render, redirect, get_object_or_404
from django.utils import timezone
from django.utils.dateparse import parse_date
from django.views.decorators.http import require_POST

from accounts.models import User
from access_requests.models import AccessRequest
from .models import VisitorCard, CardTransaction


def is_security(user):
    return user.is_authenticated and user.role == "SECURITY"


@login_required
@user_passes_test(is_security, login_url="/accounts/login/")
def card_issuance_list(request):
    now = timezone.now()

    # Transaksi yang masih aktif: kartu belum dikembalikan.
    active_transactions = CardTransaction.objects.filter(
        status__in=[
            CardTransaction.Status.BORROWED,
            CardTransaction.Status.OVERDUE,
        ]
    ).select_related(
        "student",
        "card",
        "access_request",
    ).order_by("due_at")

    # ID mahasiswa yang masih memegang kartu.
    students_with_active_cards = active_transactions.values_list(
        "student_id",
        flat=True,
    )

    # Permintaan yang sudah memiliki transaksi tidak boleh muncul
    # lagi sebagai permintaan baru untuk diberikan kartu.
    requests_with_transactions = CardTransaction.objects.values_list(
        "access_request_id",
        flat=True,
    )

    # Hanya permintaan yang benar-benar siap dilayani.
    waiting_requests = AccessRequest.objects.filter(
        status=AccessRequest.Status.APPROVED,
    ).exclude(
        id__in=requests_with_transactions,
    ).exclude(
        student_id__in=students_with_active_cards,
    ).select_related(
        "student",
    ).order_by("created_at")

    # Pisahkan transaksi yang sudah melewati batas waktu.
    overdue_transactions = active_transactions.filter(
        due_at__lt=now,
    )

    # Transaksi yang masih aktif dan belum melewati batas waktu.
    borrowed_transactions = active_transactions.filter(
        due_at__gte=now,
    )

    available_cards = VisitorCard.objects.filter(
        status=VisitorCard.Status.AVAILABLE,
    ).order_by("card_number")

    context = {
        "waiting_requests": waiting_requests,
        "borrowed_transactions": borrowed_transactions,
        "overdue_transactions": overdue_transactions,
        "available_cards": available_cards,
        "now": now,
        "waiting_count": waiting_requests.count(),
        "borrowed_count": borrowed_transactions.count(),
        "overdue_count": overdue_transactions.count(),
    }

    return render(
        request,
        "visitor_cards/issuance_list.html",
        context,
    )


@login_required
@user_passes_test(is_security, login_url="/accounts/login/")
@require_POST
def issue_card(request, request_id):
    access_request = get_object_or_404(
        AccessRequest.objects.select_related("student"),
        pk=request_id,
        status=AccessRequest.Status.APPROVED,
    )

    card_id = request.POST.get("card_id")
    due_at_input = request.POST.get("due_at")

    if not card_id or not due_at_input:
        messages.error(
            request,
            "Pilih kartu dan batas waktu pengembalian."
        )
        return redirect("card_issuance_list")

    try:
        due_at = timezone.datetime.strptime(
            due_at_input,
            "%Y-%m-%dT%H:%M"
        )
        due_at = timezone.make_aware(
            due_at,
            timezone.get_current_timezone()
        )
    except (ValueError, TypeError):
        messages.error(request, "Format batas waktu tidak valid.")
        return redirect("card_issuance_list")

    if due_at <= timezone.now():
        messages.error(
            request,
            "Batas waktu pengembalian harus di masa depan."
        )
        return redirect("card_issuance_list")

    with transaction.atomic():
        access_request = AccessRequest.objects.select_for_update().get(
            pk=request_id
        )

        if access_request.status != AccessRequest.Status.APPROVED:
            messages.error(
                request,
                "Permintaan ini tidak lagi berstatus disetujui."
            )
            return redirect("card_issuance_list")

        if CardTransaction.objects.filter(
            access_request=access_request
        ).exists():
            messages.error(
                request,
                "Permintaan ini sudah memiliki transaksi kartu."
            )
            return redirect("card_issuance_list")

        active_transaction = CardTransaction.objects.filter(
            student=access_request.student,
            status__in=[
                CardTransaction.Status.BORROWED,
                CardTransaction.Status.OVERDUE,
            ],
        ).exists()

        if active_transaction:
            messages.error(
                request,
                "Mahasiswa ini masih memiliki kartu yang belum dikembalikan. "
                "Kartu baru tidak dapat diberikan sebelum kartu sebelumnya dikembalikan."
            )
            return redirect("card_issuance_list")

        card = get_object_or_404(
            VisitorCard.objects.select_for_update(),
            pk=card_id,
            status=VisitorCard.Status.AVAILABLE,
        )

        CardTransaction.objects.create(
            student=access_request.student,
            card=card,
            access_request=access_request,
            borrowed_at=timezone.now(),
            due_at=due_at,
            status=CardTransaction.Status.BORROWED,
            approved_by=access_request.reviewed_by,
            issued_by=request.user,
        )

        card.status = VisitorCard.Status.BORROWED
        card.save(update_fields=["status"])

    messages.success(
        request,
        f"Kartu {card.card_number} berhasil diberikan kepada "
        f"{access_request.student.username}."
    )
    return redirect("card_issuance_list")

@login_required
@user_passes_test(is_security, login_url="/accounts/login/")
@require_POST
def return_card(request, transaction_id):
    with transaction.atomic():
        card_transaction = get_object_or_404(
            CardTransaction.objects.select_for_update(),
            pk=transaction_id,
            status__in=[
                CardTransaction.Status.BORROWED,
                CardTransaction.Status.OVERDUE,
            ],
        )

        card = VisitorCard.objects.select_for_update().get(
            pk=card_transaction.card_id
        )

        student = User.objects.select_for_update().get(
            pk=card_transaction.student_id
        )

        was_overdue = (
            card_transaction.status == CardTransaction.Status.OVERDUE
        )

        card_transaction.status = CardTransaction.Status.RETURNED
        card_transaction.returned_at = timezone.now()
        card_transaction.returned_to = request.user
        card_transaction.save(
            update_fields=[
                "status",
                "returned_at",
                "returned_to",
            ]
        )

        card.status = VisitorCard.Status.AVAILABLE
        card.save(update_fields=["status"])

        if student.overdue_blocked:
            other_overdue_exists = CardTransaction.objects.filter(
                student=student,
                status=CardTransaction.Status.OVERDUE,
            ).exists()

            if not other_overdue_exists:
                student.is_blocked = False
                student.overdue_blocked = False
                student.save(
                    update_fields=["is_blocked", "overdue_blocked"]
                )

    if was_overdue:
        messages.success(
            request,
            f"Kartu {card.card_number} berhasil dikembalikan. "
            "Status blokir akibat keterlambatan telah diperiksa.",
        )
    else:
        messages.success(
            request,
            f"Kartu {card.card_number} berhasil dikembalikan.",
        )

    return redirect("card_issuance_list")

@login_required
@user_passes_test(is_security, login_url="/accounts/login/")
def borrowed_cards_list(request):
    # Kartu yang masih dipinjam atau sudah melewati batas waktu
    borrowed_transactions = (
        CardTransaction.objects.filter(
            status__in=[
                CardTransaction.Status.BORROWED,
                CardTransaction.Status.OVERDUE,
            ]
        )
        .select_related(
            "student",
            "card",
            "access_request",
            "issued_by",
        )
        .order_by("due_at")
    )

    # Riwayat kartu yang sudah dikembalikan
    returned_transactions = (
        CardTransaction.objects.filter(
            status=CardTransaction.Status.RETURNED,
        )
        .select_related(
            "student",
            "card",
            "access_request",
            "issued_by",
            "returned_to",
        )
        .order_by("-returned_at")[:20]
    )

    return render(
        request,
        "visitor_cards/borrowed_cards.html",
        {
            "borrowed_transactions": borrowed_transactions,
            "returned_transactions": returned_transactions,
        },
    )

@login_required
def card_management(request):
    if request.user.role != "ADMIN":
        messages.error(
            request,
            "Kamu tidak memiliki akses ke halaman ini."
        )
        return redirect("dashboard")

    keyword = request.GET.get("q", "").strip()

    cards = VisitorCard.objects.all()

    if keyword:
        cards = cards.filter(card_number__icontains=keyword)

    context = {
        "cards": cards,
        "keyword": keyword,
        "total_cards": VisitorCard.objects.count(),
        "available_cards": VisitorCard.objects.filter(
            status=VisitorCard.Status.AVAILABLE
        ).count(),
        "borrowed_cards": VisitorCard.objects.filter(
            status=VisitorCard.Status.BORROWED
        ).count(),
        "blocked_cards": VisitorCard.objects.filter(
            status=VisitorCard.Status.BLOCKED
        ).count(),
    }

    return render(
        request,
        "visitor_cards/card_management.html",
        context,
    )

@login_required
def create_card(request):
    if request.user.role != "ADMIN":
        messages.error(
            request,
            "Kamu tidak memiliki akses ke halaman ini."
        )
        return redirect("dashboard")

    if request.method == "POST":
        card_number = request.POST.get("card_number", "").strip().upper()

        if not card_number:
            messages.error(request, "Nomor kartu wajib diisi.")

        elif len(card_number) > 20:
            messages.error(
                request,
                "Nomor kartu maksimal 20 karakter."
            )

        elif VisitorCard.objects.filter(
            card_number__iexact=card_number
        ).exists():
            messages.error(
                request,
                f"Kartu {card_number} sudah terdaftar."
            )

        else:
            VisitorCard.objects.create(
                card_number=card_number,
                status=VisitorCard.Status.AVAILABLE,
            )

            messages.success(
                request,
                f"Kartu {card_number} berhasil ditambahkan."
            )
            return redirect("card_management")

    return render(
        request,
        "visitor_cards/card_create.html",
    )

@login_required
@require_POST
def toggle_card_block(request, card_id):
    if request.user.role != "ADMIN":
        messages.error(
            request,
            "Kamu tidak memiliki akses ke tindakan ini."
        )
        return redirect("dashboard")

    with transaction.atomic():
        card = get_object_or_404(
            VisitorCard.objects.select_for_update(),
            pk=card_id,
        )

        if card.status == VisitorCard.Status.BORROWED:
            messages.error(
                request,
                "Kartu sedang dipinjam dan tidak dapat diubah "
                "statusnya dari halaman ini."
            )
            return redirect("card_management")

        if card.status == VisitorCard.Status.AVAILABLE:
            card.status = VisitorCard.Status.BLOCKED
            message = f"Kartu {card.card_number} berhasil diblokir."

        elif card.status == VisitorCard.Status.BLOCKED:
            card.status = VisitorCard.Status.AVAILABLE
            message = (
                f"Kartu {card.card_number} berhasil dibuka blokirnya."
            )

        else:
            messages.error(
                request,
                "Status kartu tidak valid."
            )
            return redirect("card_management")

        card.save(update_fields=["status"])

    messages.success(request, message)
    return redirect("card_management")

@login_required
def get_filtered_transactions(request):
    """Mengambil riwayat transaksi sesuai filter yang dipilih Admin."""
    keyword = request.GET.get("q", "").strip()
    status = request.GET.get("status", "").strip()
    start_date_input = request.GET.get("start_date", "").strip()
    end_date_input = request.GET.get("end_date", "").strip()

    transactions = CardTransaction.objects.select_related(
        "student",
        "card",
        "access_request",
        "approved_by",
        "issued_by",
        "returned_to",
    ).all()

    # Pencarian mahasiswa, NIM, dan nomor kartu.
    if keyword:
        transactions = transactions.filter(
            models.Q(student__username__icontains=keyword)
            | models.Q(student__first_name__icontains=keyword)
            | models.Q(student__last_name__icontains=keyword)
            | models.Q(student__student_id__icontains=keyword)
            | models.Q(card__card_number__icontains=keyword)
        )

    # Filter status transaksi.
    valid_statuses = [
        choice[0] for choice in CardTransaction.Status.choices
    ]

    if status in valid_statuses:
        transactions = transactions.filter(status=status)
    else:
        status = ""

    # Validasi tanggal mulai dan tanggal akhir.
    start_date = parse_date(start_date_input) if start_date_input else None
    end_date = parse_date(end_date_input) if end_date_input else None

    if start_date_input and start_date is None:
        messages.warning(request, "Format tanggal mulai tidak valid.")

    if end_date_input and end_date is None:
        messages.warning(request, "Format tanggal akhir tidak valid.")

    # Jika kedua tanggal valid tetapi urutannya salah, hasil dikosongkan.
    if start_date and end_date and start_date > end_date:
        messages.warning(
            request,
            "Tanggal mulai tidak boleh melewati tanggal akhir.",
        )
        transactions = transactions.none()
    else:
        if start_date:
            transactions = transactions.filter(
                borrowed_at__date__gte=start_date
            )

        if end_date:
            transactions = transactions.filter(
                borrowed_at__date__lte=end_date
            )

    transactions = transactions.order_by("-created_at")

    filters = {
        "keyword": keyword,
        "selected_status": status,
        "start_date": start_date_input if start_date else "",
        "end_date": end_date_input if end_date else "",
    }

    return transactions, filters


@login_required
def transaction_history(request):
    if request.user.role != "ADMIN":
        messages.error(
            request,
            "Kamu tidak memiliki akses ke halaman ini.",
        )
        return redirect("dashboard")

    transactions, filters = get_filtered_transactions(request)

    context = {
        "transactions": transactions,
        **filters,
        "total_transactions": CardTransaction.objects.count(),
        "borrowed_count": CardTransaction.objects.filter(
            status=CardTransaction.Status.BORROWED
        ).count(),
        "returned_count": CardTransaction.objects.filter(
            status=CardTransaction.Status.RETURNED
        ).count(),
        "overdue_count": CardTransaction.objects.filter(
            status=CardTransaction.Status.OVERDUE
        ).count(),
        "status_choices": CardTransaction.Status.choices,
    }

    return render(
        request,
        "visitor_cards/transaction_history.html",
        context,
    )


@login_required
def export_transaction_csv(request):
    if request.user.role != "ADMIN":
        messages.error(
            request,
            "Kamu tidak memiliki akses untuk mengekspor transaksi.",
        )
        return redirect("dashboard")

    transactions, _ = get_filtered_transactions(request)

    response = HttpResponse(
        content_type="text/csv; charset=utf-8",
    )
    response["Content-Disposition"] = (
        'attachment; filename="riwayat_transaksi.csv"'
    )

    # BOM membantu Excel membaca karakter UTF-8 dengan benar.
    response.write("\ufeff")

    writer = csv.writer(response)
    writer.writerow([
        "ID Transaksi",
        "Username",
        "Nama Mahasiswa",
        "NIM",
        "Nomor Kartu",
        "Status",
        "Waktu Pinjam",
        "Batas Pengembalian",
        "Waktu Pengembalian",
        "Disetujui Oleh",
        "Diserahkan Oleh",
        "Diterima Kembali Oleh",
    ])

    for item in transactions:
        writer.writerow([
            item.pk,
            item.student.username,
            item.student.get_full_name(),
            item.student.student_id,
            item.card.card_number,
            item.get_status_display(),
            timezone.localtime(item.borrowed_at).strftime("%Y-%m-%d %H:%M")
            if item.borrowed_at else "",
            timezone.localtime(item.due_at).strftime("%Y-%m-%d %H:%M")
            if item.due_at else "",
            timezone.localtime(item.returned_at).strftime("%Y-%m-%d %H:%M")
            if item.returned_at else "",
            item.approved_by.get_full_name() or item.approved_by.username
            if item.approved_by else "",
            item.issued_by.get_full_name() or item.issued_by.username
            if item.issued_by else "",
            item.returned_to.get_full_name() or item.returned_to.username
            if item.returned_to else "",
        ])

    return response


@login_required
def export_transaction_excel(request):
    if request.user.role != "ADMIN":
        messages.error(
            request,
            "Kamu tidak memiliki akses untuk mengekspor transaksi.",
        )
        return redirect("dashboard")

    transactions, _ = get_filtered_transactions(request)

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Riwayat Transaksi"

    worksheet.append([
        "ID Transaksi",
        "Username",
        "Nama Mahasiswa",
        "NIM",
        "Nomor Kartu",
        "Status",
        "Waktu Pinjam",
        "Batas Pengembalian",
        "Waktu Pengembalian",
        "Disetujui Oleh",
        "Diserahkan Oleh",
        "Diterima Kembali Oleh",
    ])

    for item in transactions:
        worksheet.append([
            item.pk,
            item.student.username,
            item.student.get_full_name(),
            item.student.student_id,
            item.card.card_number,
            item.get_status_display(),
            timezone.localtime(item.borrowed_at).replace(tzinfo=None)
            if item.borrowed_at else None,
            timezone.localtime(item.due_at).replace(tzinfo=None)
            if item.due_at else None,
            timezone.localtime(item.returned_at).replace(tzinfo=None)
            if item.returned_at else None,
            item.approved_by.get_full_name() or item.approved_by.username
            if item.approved_by else "",
            item.issued_by.get_full_name() or item.issued_by.username
            if item.issued_by else "",
            item.returned_to.get_full_name() or item.returned_to.username
            if item.returned_to else "",
        ])

    # Rapikan tampilan kolom.
    from openpyxl.styles import Font, PatternFill, Alignment

    for cell in worksheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(
            fill_type="solid",
            fgColor="1F4E78",
        )
        cell.alignment = Alignment(horizontal="center")

    column_widths = {
        "A": 15,
        "B": 20,
        "C": 28,
        "D": 18,
        "E": 18,
        "F": 20,
        "G": 22,
        "H": 24,
        "I": 24,
    }

    for column, width in column_widths.items():
        worksheet.column_dimensions[column].width = width

    for row in worksheet.iter_rows(min_row=2, min_col=7, max_col=9):
        for cell in row:
            cell.number_format = "dd mmm yyyy hh:mm"

    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions

    response = HttpResponse(
        content_type=(
            "application/vnd.openxmlformats-officedocument."
            "spreadsheetml.sheet"
        ),
    )
    response["Content-Disposition"] = (
        'attachment; filename="riwayat_transaksi.xlsx"'
    )

    workbook.save(response)

    return response
